"""Aplicación de consola para gestionar pedidos con openpyxl."""

from datetime import date, datetime
from pathlib import Path
from shutil import copy2
from collections import defaultdict
import re

from openpyxl import load_workbook
from openpyxl.chart import BarChart, Reference
from openpyxl.styles import Alignment, Font, PatternFill


CARPETA = Path(__file__).resolve().parent
RUTA_ORIGINAL = CARPETA / "BDD" / "BaseDatosProyecto.xlsx"
RUTA_TRABAJO = CARPETA / "BDD" / "BaseDatosProyecto_consola.xlsx"
NOMBRE_HOJA = "G. Pedidos"

CABECERAS_BASE = [
    "Cliente",
    "Artículo",
    "Fecha Pedido",
    "Fecha Entrega",
    "País Destino",
    "Ciudad destino",
    "Precio unidad",
    "Cantidad",
    "Descuento",
]

CAMPOS_ACTUALIZABLES = [
    "Precio unidad",
    "Cantidad",
    "Descuento",
]


def normalizar_cabecera(valor):
    """Elimina saltos de línea y espacios repetidos."""
    return re.sub(r"\s+", " ", str(valor)).strip()


def calcular_total(precio, cantidad, descuento):
    """Calcula precio por cantidad después de aplicar el descuento."""
    precio = float(precio)
    cantidad = int(cantidad)
    descuento = float(descuento)

    if precio < 0:
        raise ValueError("El precio no puede ser negativo.")
    if cantidad <= 0:
        raise ValueError("La cantidad debe ser mayor que cero.")
    if not 0 <= descuento <= 1:
        raise ValueError("El descuento debe estar entre 0 y 1.")

    return round(precio * cantidad * (1 - descuento), 2)


def convertir_fecha(valor, campo):
    """Convierte una fecha en formato AAAA-MM-DD."""
    if isinstance(valor, datetime):
        return valor
    if isinstance(valor, date):
        return datetime.combine(valor, datetime.min.time())
    try:
        return datetime.strptime(str(valor), "%Y-%m-%d")
    except ValueError as error:
        raise ValueError(
            f"{campo} debe utilizar el formato AAAA-MM-DD."
        ) from error


class GestorPedidos:
    """Gestiona el libro Excel y las operaciones CRUD."""

    def __init__(self):
        if not RUTA_ORIGINAL.exists():
            raise FileNotFoundError(f"No se encontró {RUTA_ORIGINAL}")

        if not RUTA_TRABAJO.exists():
            copy2(RUTA_ORIGINAL, RUTA_TRABAJO)

        self.libro = load_workbook(RUTA_TRABAJO)
        nombre = (
            NOMBRE_HOJA
            if NOMBRE_HOJA in self.libro.sheetnames
            else self.libro.sheetnames[0]
        )
        self.hoja = self.libro[nombre]
        self._preparar_hoja()

    def _indices(self):
        return {
            normalizar_cabecera(self.hoja.cell(1, columna).value): columna
            for columna in range(1, self.hoja.max_column + 1)
        }

    def _preparar_hoja(self):
        for columna in range(1, self.hoja.max_column + 1):
            celda = self.hoja.cell(1, columna)
            celda.value = normalizar_cabecera(celda.value)

        indices = self._indices()
        faltantes = [
            cabecera for cabecera in CABECERAS_BASE
            if cabecera not in indices
        ]
        if faltantes:
            raise ValueError(f"Faltan cabeceras: {faltantes}")

        if "Total" not in indices:
            self.hoja.cell(1, self.hoja.max_column + 1, "Total")

        self.recalcular_todos()
        self.guardar()

    def guardar(self):
        self.libro.save(RUTA_TRABAJO)

    def recalcular_todos(self):
        indices = self._indices()
        for fila in range(2, self.hoja.max_row + 1):
            total = calcular_total(
                self.hoja.cell(fila, indices["Precio unidad"]).value,
                self.hoja.cell(fila, indices["Cantidad"]).value,
                self.hoja.cell(fila, indices["Descuento"]).value,
            )
            celda = self.hoja.cell(fila, indices["Total"])
            celda.value = total
            celda.number_format = "#,##0.00"

    def leer(self, fila_excel):
        fila_excel = int(fila_excel)
        if fila_excel < 2 or fila_excel > self.hoja.max_row:
            raise KeyError(f"No existe la fila Excel {fila_excel}.")

        nombres = list(self._indices())
        valores = [
            self.hoja.cell(fila_excel, columna).value
            for columna in range(1, self.hoja.max_column + 1)
        ]
        registro = {"Fila Excel": fila_excel}
        registro.update(dict(zip(nombres, valores)))
        return registro

    def listar(self):
        return [
            self.leer(fila)
            for fila in range(2, self.hoja.max_row + 1)
        ]

    def crear(self, pedido):
        fecha_pedido = convertir_fecha(
            pedido["Fecha Pedido"], "Fecha Pedido"
        )
        fecha_entrega = convertir_fecha(
            pedido["Fecha Entrega"], "Fecha Entrega"
        )
        if fecha_entrega < fecha_pedido:
            raise ValueError(
                "Fecha Entrega no puede ser anterior a Fecha Pedido."
            )

        valores = dict(pedido)
        valores["Fecha Pedido"] = fecha_pedido
        valores["Fecha Entrega"] = fecha_entrega
        valores["Precio unidad"] = float(valores["Precio unidad"])
        valores["Cantidad"] = int(valores["Cantidad"])
        valores["Descuento"] = float(valores["Descuento"])
        valores["Total"] = calcular_total(
            valores["Precio unidad"],
            valores["Cantidad"],
            valores["Descuento"],
        )

        for campo in [
            "Cliente", "Artículo", "País Destino", "Ciudad destino"
        ]:
            if not str(valores[campo]).strip():
                raise ValueError(f"{campo} es obligatorio.")

        fila = self.hoja.max_row + 1
        for campo, columna in self._indices().items():
            self.hoja.cell(fila, columna, valores.get(campo))

        self.guardar()
        return self.leer(fila)

    def actualizar(self, fila_excel, campo, nuevo_valor):
        if campo not in CAMPOS_ACTUALIZABLES:
            raise KeyError(
                "Solo se pueden actualizar Precio unidad, "
                "Cantidad y Descuento."
            )

        fila_excel = int(fila_excel)
        self.leer(fila_excel)
        indices = self._indices()

        if campo == "Cantidad":
            nuevo_valor = int(nuevo_valor)
        else:
            nuevo_valor = float(nuevo_valor)

        precio = (
            nuevo_valor
            if campo == "Precio unidad"
            else self.hoja.cell(fila_excel, indices["Precio unidad"]).value
        )
        cantidad = (
            nuevo_valor
            if campo == "Cantidad"
            else self.hoja.cell(fila_excel, indices["Cantidad"]).value
        )
        descuento = (
            nuevo_valor
            if campo == "Descuento"
            else self.hoja.cell(fila_excel, indices["Descuento"]).value
        )
        total = calcular_total(precio, cantidad, descuento)

        self.hoja.cell(fila_excel, indices[campo], nuevo_valor)
        self.hoja.cell(fila_excel, indices["Total"], total)
        self.guardar()
        return self.leer(fila_excel)

    def eliminar(self, fila_excel):
        eliminado = self.leer(fila_excel)
        self.hoja.delete_rows(int(fila_excel), 1)
        self.guardar()
        return eliminado

    def filtrar(self, cliente="", articulo="", pais="", ciudad=""):
        textos = {
            "Cliente": cliente.strip().lower(),
            "Artículo": articulo.strip().lower(),
            "País Destino": pais.strip().lower(),
            "Ciudad destino": ciudad.strip().lower(),
        }
        resultado = []
        for registro in self.listar():
            coincide = all(
                not texto
                or texto in str(registro[campo]).lower()
                for campo, texto in textos.items()
            )
            if coincide:
                resultado.append(registro)
        return resultado

    def generar_resumen_y_grafico(self):
        """Crea una hoja Resumen y un gráfico del total por país."""
        acumulados = defaultdict(
            lambda: {"Pedidos": 0, "Cantidad": 0, "Total": 0.0}
        )

        for registro in self.listar():
            pais = registro["País Destino"]
            acumulados[pais]["Pedidos"] += 1
            acumulados[pais]["Cantidad"] += registro["Cantidad"]
            acumulados[pais]["Total"] += registro["Total"]

        resumen = [
            {
                "País Destino": pais,
                "Pedidos": valores["Pedidos"],
                "Cantidad": valores["Cantidad"],
                "Total": round(valores["Total"], 2),
            }
            for pais, valores in acumulados.items()
        ]
        resumen.sort(key=lambda registro: registro["Total"], reverse=True)

        if "Resumen" in self.libro.sheetnames:
            self.libro.remove(self.libro["Resumen"])

        hoja_resumen = self.libro.create_sheet("Resumen", 0)
        hoja_resumen.append([
            "País Destino",
            "Pedidos",
            "Cantidad",
            "Total",
        ])

        for registro in resumen:
            hoja_resumen.append([
                registro["País Destino"],
                registro["Pedidos"],
                registro["Cantidad"],
                registro["Total"],
            ])

        for celda in hoja_resumen[1]:
            celda.font = Font(bold=True, color="FFFFFF")
            celda.fill = PatternFill("solid", fgColor="1F4E78")
            celda.alignment = Alignment(horizontal="center")

        for fila in range(2, hoja_resumen.max_row + 1):
            hoja_resumen.cell(fila, 4).number_format = "#,##0.00"

        hoja_resumen.column_dimensions["A"].width = 22
        hoja_resumen.column_dimensions["B"].width = 12
        hoja_resumen.column_dimensions["C"].width = 14
        hoja_resumen.column_dimensions["D"].width = 16
        hoja_resumen.sheet_properties.pageSetUpPr.fitToPage = True
        hoja_resumen.page_setup.orientation = "landscape"
        hoja_resumen.page_setup.fitToWidth = 1
        hoja_resumen.page_setup.fitToHeight = 1
        hoja_resumen.print_area = f"A1:N{max(24, hoja_resumen.max_row)}"

        cantidad_paises = min(10, len(resumen))
        if cantidad_paises:
            grafico = BarChart()
            grafico.title = "Total por país"
            grafico.y_axis.title = "Total"
            grafico.x_axis.title = "País"

            datos = Reference(
                hoja_resumen,
                min_col=4,
                min_row=1,
                max_row=cantidad_paises + 1,
            )
            categorias = Reference(
                hoja_resumen,
                min_col=1,
                min_row=2,
                max_row=cantidad_paises + 1,
            )
            grafico.add_data(datos, titles_from_data=True)
            grafico.set_categories(categorias)
            grafico.height = 8
            grafico.width = 15
            hoja_resumen.add_chart(grafico, "F2")

        self.guardar()
        return resumen


def mostrar_registros(registros, limite=20):
    registros = list(registros)
    if not registros:
        print("No se encontraron registros.")
        return

    columnas = [
        "Fila Excel", "Cliente", "Artículo", "País Destino",
        "Cantidad", "Descuento", "Total",
    ]
    anchos = {
        columna: min(
            35,
            max(
                len(columna),
                *(len(str(r.get(columna, ""))) for r in registros[:limite]),
            ),
        )
        for columna in columnas
    }

    encabezado = " | ".join(
        f"{columna:<{anchos[columna]}}" for columna in columnas
    )
    print(encabezado)
    print("-" * len(encabezado))
    for registro in registros[:limite]:
        print(" | ".join(
            f"{str(registro.get(columna, '')):<{anchos[columna]}}"
            for columna in columnas
        ))
    print(f"\nMostrando {min(limite, len(registros))} de {len(registros)}.")


def pedir_texto(nombre):
    while True:
        valor = input(f"{nombre}: ").strip()
        if valor:
            return valor
        print(f"Error: {nombre} es obligatorio.")


def pedir_fecha(nombre):
    while True:
        valor = input(f"{nombre} (AAAA-MM-DD): ").strip()
        try:
            fecha = convertir_fecha(valor, nombre)
            print(f"Fecha válida: {fecha:%d/%m/%Y}")
            return fecha
        except ValueError as error:
            print(f"Error: {error}")


def pedir_precio():
    while True:
        try:
            valor = float(input("Precio unidad: ").strip())
            calcular_total(valor, 1, 0)
            return valor
        except ValueError as error:
            print(f"Error: {error}")


def pedir_cantidad():
    while True:
        try:
            valor = int(input("Cantidad: ").strip())
            calcular_total(0, valor, 0)
            return valor
        except ValueError as error:
            print(f"Error: {error}")


def pedir_descuento():
    while True:
        try:
            valor = float(input("Descuento entre 0 y 1: ").strip())
            calcular_total(0, 1, valor)
            return valor
        except ValueError as error:
            print(f"Error: {error}")


def buscar_fila_para_actualizar(gestor):
    """Busca pedidos y permite seleccionar una fila de los resultados."""
    print("\nBÚSQUEDA PREVIA DEL PEDIDO")
    print("Deja vacíos los filtros que no quieras utilizar.")
    resultados = gestor.filtrar(
        cliente=input("Cliente: "),
        articulo=input("Artículo: "),
        pais=input("País: "),
        ciudad=input("Ciudad: "),
    )

    if not resultados:
        print("No se encontraron pedidos con esos filtros.")
        return None

    mostrar_registros(resultados, limite=50)
    filas_permitidas = {
        registro["Fila Excel"] for registro in resultados
    }

    while True:
        valor = input(
            "Escribe la fila Excel que quieres actualizar "
            "o pulsa Enter para cancelar: "
        ).strip()
        if valor == "":
            print("Actualización cancelada.")
            return None

        try:
            fila = int(valor)
        except ValueError:
            print("La fila debe ser un número entero.")
            continue

        if fila not in filas_permitidas:
            print("La fila elegida no pertenece a los resultados mostrados.")
            continue

        print("\nPedido seleccionado:")
        mostrar_registros([gestor.leer(fila)], limite=1)
        return fila


def menu(gestor):
    while True:
        print("\n" + "=" * 58)
        print("        CRUD DE PEDIDOS - OPENPYXL")
        print("=" * 58)
        print("1. Crear pedido")
        print("2. Consultar pedido")
        print("3. Filtrar pedidos")
        print("4. Actualizar precio, cantidad o descuento")
        print("5. Eliminar pedido")
        print("6. Mostrar los últimos registros")
        print("7. Guardar")
        print("8. Generar resumen, gráfico y salir")

        opcion = input("Selecciona una opción (1-8): ").strip()

        try:
            if opcion == "1":
                fecha_pedido = pedir_fecha("Fecha Pedido")
                while True:
                    fecha_entrega = pedir_fecha("Fecha Entrega")
                    if fecha_entrega >= fecha_pedido:
                        break
                    print("Error: la entrega no puede ser anterior al pedido.")

                pedido = {
                    "Cliente": pedir_texto("Cliente"),
                    "Artículo": pedir_texto("Artículo"),
                    "Fecha Pedido": fecha_pedido,
                    "Fecha Entrega": fecha_entrega,
                    "País Destino": pedir_texto("País Destino"),
                    "Ciudad destino": pedir_texto("Ciudad destino"),
                    "Precio unidad": pedir_precio(),
                    "Cantidad": pedir_cantidad(),
                    "Descuento": pedir_descuento(),
                }
                mostrar_registros([gestor.crear(pedido)])

            elif opcion == "2":
                fila = int(input("Fila Excel: "))
                mostrar_registros([gestor.leer(fila)])

            elif opcion == "3":
                print("Deja un campo vacío para no aplicar ese filtro.")
                resultado = gestor.filtrar(
                    cliente=input("Cliente: "),
                    articulo=input("Artículo: "),
                    pais=input("País: "),
                    ciudad=input("Ciudad: "),
                )
                mostrar_registros(resultado)

            elif opcion == "4":
                fila = buscar_fila_para_actualizar(gestor)
                if fila is None:
                    continue
                print("1. Precio unidad")
                print("2. Cantidad")
                print("3. Descuento")
                seleccion = input("Campo (1-3): ").strip()
                campos = {
                    "1": ("Precio unidad", pedir_precio),
                    "2": ("Cantidad", pedir_cantidad),
                    "3": ("Descuento", pedir_descuento),
                }
                if seleccion not in campos:
                    print("Campo incorrecto.")
                    continue
                campo, solicitar = campos[seleccion]
                mostrar_registros([
                    gestor.actualizar(fila, campo, solicitar())
                ])

            elif opcion == "5":
                fila = int(input("Fila Excel: "))
                mostrar_registros([gestor.leer(fila)])
                confirmar = input(
                    "Escribe ELIMINAR para confirmar: "
                ).strip()
                if confirmar == "ELIMINAR":
                    gestor.eliminar(fila)
                    print("Pedido eliminado.")
                else:
                    print("Eliminación cancelada.")

            elif opcion == "6":
                mostrar_registros(gestor.listar()[-20:])

            elif opcion == "7":
                gestor.guardar()
                print(f"Guardado en {RUTA_TRABAJO}")

            elif opcion == "8":
                resumen = gestor.generar_resumen_y_grafico()
                print(
                    f"Resumen generado para {len(resumen)} países "
                    "en la hoja 'Resumen'."
                )
                print(f"Gráfico guardado en {RUTA_TRABAJO}")
                print("Aplicación finalizada.")
                break

            else:
                print("Opción incorrecta.")

        except (ValueError, KeyError) as error:
            print(f"Error: {error}")


def main():
    gestor = GestorPedidos()
    print(f"Archivo de trabajo: {RUTA_TRABAJO}")
    menu(gestor)


if __name__ == "__main__":
    main()
