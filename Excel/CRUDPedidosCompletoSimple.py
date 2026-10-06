"""CRUD de pedidos con openpyxl.
Mantiene las funciones del CRUD completo, pero con estructura procedural sencilla.
"""

from datetime import date, datetime
from pathlib import Path
from shutil import copy2
from collections import defaultdict
import re

from openpyxl import load_workbook
from openpyxl.chart import BarChart, Reference
from openpyxl.styles import Alignment, Font, PatternFill


CARPETA = Path(__file__).resolve().parent
RUTA_ORIGINAL = CARPETA / "BDD" / "BaseDatosProyecto.xlsx"
RUTA_TRABAJO = CARPETA / "BDD" / "BaseDatosProyecto_consola.xlsx"
NOMBRE_HOJA = "G. Pedidos"

CABECERAS_BASE = [
    "Cliente",
    "Artículo",
    "Fecha Pedido",
    "Fecha Entrega",
    "País Destino",
    "Ciudad destino",
    "Precio unidad",
    "Cantidad",
    "Descuento",
]

CAMPOS_ACTUALIZABLES = [
    "Precio unidad",
    "Cantidad",
    "Descuento",
]


# ==========================================
# 1. PREPARAR EXCEL
# ==========================================

def normalizar_cabecera(valor):
    """Elimina saltos de línea y espacios repetidos."""
    return re.sub(r"\s+", " ", str(valor)).strip()


def calcular_total(precio, cantidad, descuento):
    """Calcula el total aplicando el descuento."""
    precio = float(precio)
    cantidad = int(cantidad)
    descuento = float(descuento)

    if precio < 0:
        raise ValueError("El precio no puede ser negativo.")
    if cantidad <= 0:
        raise ValueError("La cantidad debe ser mayor que cero.")
    if not 0 <= descuento <= 1:
        raise ValueError("El descuento debe estar entre 0 y 1.")

    return round(precio * cantidad * (1 - descuento), 2)


def convertir_fecha(valor, campo):
    """Convierte una fecha a datetime."""
    if isinstance(valor, datetime):
        return valor
    if isinstance(valor, date):
        return datetime.combine(valor, datetime.min.time())

    try:
        return datetime.strptime(str(valor), "%Y-%m-%d")
    except ValueError:
        raise ValueError(f"{campo} debe utilizar el formato AAAA-MM-DD.")


if not RUTA_ORIGINAL.exists():
    raise FileNotFoundError(f"No se encontró {RUTA_ORIGINAL}")

if not RUTA_TRABAJO.exists():
    copy2(RUTA_ORIGINAL, RUTA_TRABAJO)

libro = load_workbook(RUTA_TRABAJO)

if NOMBRE_HOJA in libro.sheetnames:
    hoja = libro[NOMBRE_HOJA]
else:
    hoja = libro[libro.sheetnames[0]]


def indices():
    """Devuelve cabecera: número de columna."""
    return {
        normalizar_cabecera(hoja.cell(1, columna).value): columna
        for columna in range(1, hoja.max_column + 1)
    }


def guardar():
    """Guarda el libro Excel."""
    libro.save(RUTA_TRABAJO)


def preparar_hoja():
    """Limpia cabeceras y crea la columna Total si falta."""
    for columna in range(1, hoja.max_column + 1):
        celda = hoja.cell(1, columna)
        celda.value = normalizar_cabecera(celda.value)

    columnas = indices()

    faltantes = [
        cabecera
        for cabecera in CABECERAS_BASE
        if cabecera not in columnas
    ]

    if faltantes:
        raise ValueError(f"Faltan cabeceras: {faltantes}")

    if "Total" not in columnas:
        hoja.cell(1, hoja.max_column + 1, "Total")

    recalcular_todos()
    guardar()


# ==========================================
# 2. OPERACIONES CRUD
# ==========================================

def recalcular_todos():
    """Recalcula el Total de todos los pedidos."""
    columnas = indices()

    for fila in range(2, hoja.max_row + 1):
        total = calcular_total(
            hoja.cell(fila, columnas["Precio unidad"]).value,
            hoja.cell(fila, columnas["Cantidad"]).value,
            hoja.cell(fila, columnas["Descuento"]).value,
        )

        celda = hoja.cell(fila, columnas["Total"])
        celda.value = total
        celda.number_format = "#,##0.00"


def leer(fila_excel):
    """Lee una fila de Excel y devuelve un diccionario."""
    fila_excel = int(fila_excel)

    if fila_excel < 2 or fila_excel > hoja.max_row:
        raise KeyError(f"No existe la fila Excel {fila_excel}.")

    columnas = indices()

    registro = {"Fila Excel": fila_excel}

    for nombre, columna in columnas.items():
        registro[nombre] = hoja.cell(fila_excel, columna).value

    return registro


def listar():
    """Devuelve todos los pedidos."""
    registros = []

    for fila in range(2, hoja.max_row + 1):
        registros.append(leer(fila))

    return registros


def crear(pedido):
    """Crea un pedido nuevo."""
    fecha_pedido = convertir_fecha(
        pedido["Fecha Pedido"],
        "Fecha Pedido"
    )

    fecha_entrega = convertir_fecha(
        pedido["Fecha Entrega"],
        "Fecha Entrega"
    )

    if fecha_entrega < fecha_pedido:
        raise ValueError(
            "Fecha Entrega no puede ser anterior a Fecha Pedido."
        )

    for campo in [
        "Cliente",
        "Artículo",
        "País Destino",
        "Ciudad destino"
    ]:
        if not str(pedido[campo]).strip():
            raise ValueError(f"{campo} es obligatorio.")

    valores = dict(pedido)

    valores["Fecha Pedido"] = fecha_pedido
    valores["Fecha Entrega"] = fecha_entrega
    valores["Precio unidad"] = float(valores["Precio unidad"])
    valores["Cantidad"] = int(valores["Cantidad"])
    valores["Descuento"] = float(valores["Descuento"])

    valores["Total"] = calcular_total(
        valores["Precio unidad"],
        valores["Cantidad"],
        valores["Descuento"],
    )

    fila = hoja.max_row + 1

    for campo, columna in indices().items():
        hoja.cell(
            fila,
            columna,
            valores.get(campo)
        )

    guardar()

    return leer(fila)


def actualizar(fila_excel, campo, nuevo_valor):
    """Actualiza precio, cantidad o descuento."""
    if campo not in CAMPOS_ACTUALIZABLES:
        raise KeyError(
            "Solo se pueden actualizar Precio unidad, "
            "Cantidad y Descuento."
        )

    fila_excel = int(fila_excel)

    # Verifica que la fila exista.
    leer(fila_excel)

    columnas = indices()

    if campo == "Cantidad":
        nuevo_valor = int(nuevo_valor)
    else:
        nuevo_valor = float(nuevo_valor)

    precio = hoja.cell(
        fila_excel,
        columnas["Precio unidad"]
    ).value

    cantidad = hoja.cell(
        fila_excel,
        columnas["Cantidad"]
    ).value

    descuento = hoja.cell(
        fila_excel,
        columnas["Descuento"]
    ).value

    if campo == "Precio unidad":
        precio = nuevo_valor
    elif campo == "Cantidad":
        cantidad = nuevo_valor
    elif campo == "Descuento":
        descuento = nuevo_valor

    total = calcular_total(
        precio,
        cantidad,
        descuento
    )

    hoja.cell(
        fila_excel,
        columnas[campo],
        nuevo_valor
    )

    hoja.cell(
        fila_excel,
        columnas["Total"],
        total
    )

    guardar()

    return leer(fila_excel)


def eliminar(fila_excel):
    """Elimina una fila de Excel."""
    eliminado = leer(fila_excel)

    hoja.delete_rows(
        int(fila_excel),
        1
    )

    guardar()

    return eliminado


def filtrar(cliente="", articulo="", pais="", ciudad=""):
    """Filtra pedidos por varios campos."""
    cliente = cliente.strip().lower()
    articulo = articulo.strip().lower()
    pais = pais.strip().lower()
    ciudad = ciudad.strip().lower()

    resultado = []

    for registro in listar():
        if cliente and cliente not in str(registro["Cliente"]).lower():
            continue

        if articulo and articulo not in str(registro["Artículo"]).lower():
            continue

        if pais and pais not in str(registro["País Destino"]).lower():
            continue

        if ciudad and ciudad not in str(registro["Ciudad destino"]).lower():
            continue

        resultado.append(registro)

    return resultado


# ==========================================
# 3. MOSTRAR DATOS
# ==========================================

def mostrar_registros(registros, limite=20):
    """Muestra registros en forma de tabla."""
    registros = list(registros)

    if not registros:
        print("No se encontraron registros.")
        return

    columnas = [
        "Fila Excel",
        "Cliente",
        "Artículo",
        "País Destino",
        "Cantidad",
        "Descuento",
        "Total",
    ]

    print()
    print(
        f"{'Fila':<6}"
        f"{'Cliente':<25}"
        f"{'Artículo':<28}"
        f"{'País':<18}"
        f"{'Cantidad':>10}"
        f"{'Descuento':>12}"
        f"{'Total':>14}"
    )

    print("-" * 113)

    for registro in registros[:limite]:
        print(
            f"{registro['Fila Excel']:<6}"
            f"{str(registro['Cliente'])[:23]:<25}"
            f"{str(registro['Artículo'])[:26]:<28}"
            f"{str(registro['País Destino'])[:16]:<18}"
            f"{str(registro['Cantidad']):>10}"
            f"{str(registro['Descuento']):>12}"
            f"{str(registro['Total']):>14}"
        )

    print(
        f"\nMostrando "
        f"{min(limite, len(registros))} "
        f"de {len(registros)}."
    )


# ==========================================
# 4. VALIDACIONES DE ENTRADA
# ==========================================

def pedir_texto(nombre):
    while True:
        valor = input(f"{nombre}: ").strip()

        if valor:
            return valor

        print(f"Error: {nombre} es obligatorio.")


def pedir_fecha(nombre):
    while True:
        valor = input(
            f"{nombre} (AAAA-MM-DD): "
        ).strip()

        try:
            return convertir_fecha(
                valor,
                nombre
            )
        except ValueError as error:
            print(f"Error: {error}")


def pedir_precio():
    while True:
        try:
            valor = float(
                input("Precio unidad: ").strip()
            )

            calcular_total(
                valor,
                1,
                0
            )

            return valor

        except ValueError as error:
            print(f"Error: {error}")


def pedir_cantidad():
    while True:
        try:
            valor = int(
                input("Cantidad: ").strip()
            )

            calcular_total(
                0,
                valor,
                0
            )

            return valor

        except ValueError as error:
            print(f"Error: {error}")


def pedir_descuento():
    while True:
        try:
            valor = float(
                input(
                    "Descuento entre 0 y 1: "
                ).strip()
            )

            calcular_total(
                0,
                1,
                valor
            )

            return valor

        except ValueError as error:
            print(f"Error: {error}")


# ==========================================
# 5. BÚSQUEDA PARA ACTUALIZAR
# ==========================================

def buscar_fila_para_actualizar():
    """Busca un pedido antes de modificarlo."""
    print("\nBÚSQUEDA PREVIA DEL PEDIDO")
    print(
        "Deja vacío un filtro "
        "si no quieres utilizarlo."
    )

    resultados = filtrar(
        cliente=input("Cliente: "),
        articulo=input("Artículo: "),
        pais=input("País: "),
        ciudad=input("Ciudad: "),
    )

    if not resultados:
        print("No se encontraron pedidos.")
        return None

    mostrar_registros(
        resultados,
        limite=50
    )

    filas_permitidas = [
        registro["Fila Excel"]
        for registro in resultados
    ]

    while True:
        valor = input(
            "Fila Excel a actualizar "
            "(Enter para cancelar): "
        ).strip()

        if valor == "":
            return None

        try:
            fila = int(valor)
        except ValueError:
            print(
                "La fila debe ser un número."
            )
            continue

        if fila not in filas_permitidas:
            print(
                "La fila no pertenece "
                "a los resultados."
            )
            continue

        return fila


# ==========================================
# 6. RESUMEN Y GRÁFICO DE EXCEL
# ==========================================

def generar_resumen_y_grafico():
    """Crea hoja Resumen y gráfico de barras del Total por país."""

    acumulados = defaultdict(
        lambda: {
            "Pedidos": 0,
            "Cantidad": 0,
            "Total": 0.0
        }
    )

    # 1. Agrupar los datos por país.
    for registro in listar():
        pais = registro["País Destino"]

        acumulados[pais]["Pedidos"] += 1
        acumulados[pais]["Cantidad"] += registro["Cantidad"]
        acumulados[pais]["Total"] += registro["Total"]

    # 2. Convertir el resumen a una lista.
    resumen = []

    for pais, valores in acumulados.items():
        resumen.append({
            "País Destino": pais,
            "Pedidos": valores["Pedidos"],
            "Cantidad": valores["Cantidad"],
            "Total": round(
                valores["Total"],
                2
            ),
        })

    resumen.sort(
        key=lambda registro: registro["Total"],
        reverse=True
    )

    # 3. Borrar la hoja anterior si ya existe.
    if "Resumen" in libro.sheetnames:
        libro.remove(
            libro["Resumen"]
        )

    # 4. Crear una hoja nueva.
    hoja_resumen = libro.create_sheet(
        "Resumen",
        0
    )

    # 5. Escribir cabeceras.
    hoja_resumen.append([
        "País Destino",
        "Pedidos",
        "Cantidad",
        "Total",
    ])

    # 6. Escribir los datos agrupados.
    for registro in resumen:
        hoja_resumen.append([
            registro["País Destino"],
            registro["Pedidos"],
            registro["Cantidad"],
            registro["Total"],
        ])

    # 7. Formato sencillo de la cabecera.
    for celda in hoja_resumen[1]:
        celda.font = Font(
            bold=True,
            color="FFFFFF"
        )

        celda.fill = PatternFill(
            "solid",
            fgColor="1F4E78"
        )

        celda.alignment = Alignment(
            horizontal="center"
        )

    # 8. Formato numérico del total.
    for fila in range(
        2,
        hoja_resumen.max_row + 1
    ):
        hoja_resumen.cell(
            fila,
            4
        ).number_format = "#,##0.00"

    # 9. Crear gráfico con los 10 países que más facturan.
    cantidad_paises = min(
        10,
        len(resumen)
    )

    if cantidad_paises > 0:
        grafico = BarChart()

        grafico.title = "Total por país"
        grafico.y_axis.title = "Total"
        grafico.x_axis.title = "País"

        # Valores: columna D = Total.
        datos = Reference(
            hoja_resumen,
            min_col=4,
            min_row=1,
            max_row=cantidad_paises + 1,
        )

        # Categorías: columna A = País.
        categorias = Reference(
            hoja_resumen,
            min_col=1,
            min_row=2,
            max_row=cantidad_paises + 1,
        )

        grafico.add_data(
            datos,
            titles_from_data=True
        )

        grafico.set_categories(
            categorias
        )

        grafico.height = 8
        grafico.width = 15

        # Coloca el gráfico a partir de F2.
        hoja_resumen.add_chart(
            grafico,
            "F2"
        )

    guardar()

    return resumen


# ==========================================
# 7. MENÚ
# ==========================================

def menu():
    while True:
        print("\n" + "=" * 58)
        print("        CRUD DE PEDIDOS - OPENPYXL")
        print("=" * 58)

        print("1. Crear pedido")
        print("2. Consultar pedido")
        print("3. Filtrar pedidos")
        print(
            "4. Actualizar precio, "
            "cantidad o descuento"
        )
        print("5. Eliminar pedido")
        print("6. Mostrar últimos registros")
        print("7. Guardar")
        print(
            "8. Generar resumen, "
            "gráfico y salir"
        )

        opcion = input(
            "Selecciona una opción (1-8): "
        ).strip()

        try:
            if opcion == "1":
                fecha_pedido = pedir_fecha(
                    "Fecha Pedido"
                )

                while True:
                    fecha_entrega = pedir_fecha(
                        "Fecha Entrega"
                    )

                    if fecha_entrega >= fecha_pedido:
                        break

                    print(
                        "Error: la entrega no puede "
                        "ser anterior al pedido."
                    )

                pedido = {
                    "Cliente": pedir_texto(
                        "Cliente"
                    ),
                    "Artículo": pedir_texto(
                        "Artículo"
                    ),
                    "Fecha Pedido": fecha_pedido,
                    "Fecha Entrega": fecha_entrega,
                    "País Destino": pedir_texto(
                        "País Destino"
                    ),
                    "Ciudad destino": pedir_texto(
                        "Ciudad destino"
                    ),
                    "Precio unidad": pedir_precio(),
                    "Cantidad": pedir_cantidad(),
                    "Descuento": pedir_descuento(),
                }

                mostrar_registros([
                    crear(pedido)
                ])

            elif opcion == "2":
                fila = int(
                    input("Fila Excel: ")
                )

                mostrar_registros([
                    leer(fila)
                ])

            elif opcion == "3":
                print(
                    "Deja vacío un filtro "
                    "si no quieres utilizarlo."
                )

                resultado = filtrar(
                    cliente=input("Cliente: "),
                    articulo=input("Artículo: "),
                    pais=input("País: "),
                    ciudad=input("Ciudad: "),
                )

                mostrar_registros(
                    resultado
                )

            elif opcion == "4":
                fila = buscar_fila_para_actualizar()

                if fila is None:
                    continue

                print("1. Precio unidad")
                print("2. Cantidad")
                print("3. Descuento")

                seleccion = input(
                    "Campo (1-3): "
                ).strip()

                campos = {
                    "1": (
                        "Precio unidad",
                        pedir_precio
                    ),
                    "2": (
                        "Cantidad",
                        pedir_cantidad
                    ),
                    "3": (
                        "Descuento",
                        pedir_descuento
                    ),
                }

                if seleccion not in campos:
                    print("Campo incorrecto.")
                    continue

                campo, pedir_valor = (
                    campos[seleccion]
                )

                mostrar_registros([
                    actualizar(
                        fila,
                        campo,
                        pedir_valor()
                    )
                ])

            elif opcion == "5":
                fila = int(
                    input("Fila Excel: ")
                )

                mostrar_registros([
                    leer(fila)
                ])

                confirmar = input(
                    "Escribe ELIMINAR "
                    "para confirmar: "
                ).strip()

                if confirmar == "ELIMINAR":
                    eliminar(fila)
                    print("Pedido eliminado.")
                else:
                    print(
                        "Eliminación cancelada."
                    )

            elif opcion == "6":
                mostrar_registros(
                    listar()[-20:]
                )

            elif opcion == "7":
                guardar()
                print(
                    f"Guardado en "
                    f"{RUTA_TRABAJO}"
                )

            elif opcion == "8":
                resumen = (
                    generar_resumen_y_grafico()
                )

                print(
                    f"Resumen generado para "
                    f"{len(resumen)} países."
                )

                print(
                    f"Gráfico guardado en "
                    f"{RUTA_TRABAJO}"
                )

                break

            else:
                print("Opción incorrecta.")

        except (ValueError, KeyError) as error:
            print(f"Error: {error}")


def main():
    preparar_hoja()

    print(
        f"Archivo de trabajo: "
        f"{RUTA_TRABAJO}"
    )

    menu()


if __name__ == "__main__":
    main()
